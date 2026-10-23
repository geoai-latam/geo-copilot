"""Tablas de un archivo publicado, Excel o CSV (Terranova F5): anexos del DANE y directorios de la DIAN.

- El archivo vigente se descubre en su página índice (`pagina` + `enlace_con`): la DIAN pone la fecha
  en el nombre de su archivo diario y el DANE el mes en el de su anexo.
- `hoja` y `fila_encabezado`: los anexos del DANE traen títulos arriba y el encabezado en la fila 9.
- `rellenar`: las primeras columnas combinadas («Nacional» una vez, para varias filas) se rellenan
  hacia abajo.
- `largo_desde`: una tabla ancha (un año por columna) pasa a una fila por (identificadores, columna,
  valor). Así un año nuevo no rompe la lista blanca de quien la consume.
- `filtros`: igualdad, sin distinguir mayúsculas, sobre las columnas.
"""

from __future__ import annotations

import csv
import io
from typing import Any

MAX_FILAS = 5000
MAX_BYTES = 25_000_000


def _texto(valor: Any) -> str:
    return "" if valor is None else " ".join(str(valor).split())


def _encabezados(fila: tuple[Any, ...]) -> list[str]:
    nombres, vistos = [], set()
    for i, celda in enumerate(fila):
        nombre = _texto(celda) or f"columna_{i + 1}"
        while nombre in vistos:
            nombre += "_"
        vistos.add(nombre)
        nombres.append(nombre)
    return nombres


def filas_crudas(contenido: bytes, nombre: str, hoja: str | None) -> tuple[list[tuple[Any, ...]], list[str]]:
    """(filas, hojas) de un .xlsx o un .csv."""
    if len(contenido) > MAX_BYTES:
        raise ValueError(f"el archivo pesa {len(contenido)} bytes: más de {MAX_BYTES}")
    if nombre.lower().split("?")[0].endswith(".csv"):
        texto = contenido.decode("utf-8-sig", errors="replace")
        separador = ";" if texto[:2000].count(";") > texto[:2000].count(",") else ","
        return [tuple(f) for f in csv.reader(io.StringIO(texto), delimiter=separador)], []
    import openpyxl  # solo para Excel; read_only no carga el archivo entero en memoria

    libro = openpyxl.load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    hojas = list(libro.sheetnames)
    if hoja is not None and hoja not in hojas:
        # algunos anexos traen espacios al final del nombre de la hoja («IPM_Sexo »)
        hoja = next((h for h in hojas if h.strip() == hoja.strip()), None)
        if hoja is None:
            raise ValueError(f"no hay hoja con ese nombre; las hojas son: {hojas}")
    filas = list(libro[hoja or hojas[0]].iter_rows(values_only=True))
    libro.close()
    return filas, hojas


def _arriba(fila: tuple[Any, ...], ancho: int) -> list[str]:
    """La fila superior de un encabezado de dos, rellena hacia la derecha (un año combinado sobre
    sus 3 o 12 columnas)."""
    salida, previo = [], ""
    for i in range(ancho):
        actual = _texto(fila[i]) if i < len(fila) else ""
        previo = actual or previo
        salida.append(previo)
    return salida


def tabla(filas: list[tuple[Any, ...]], fila_encabezado: int = 1, rellenar: int = 0,
          largo_desde: int | None = None, filas_encabezado: int = 1) -> list[dict[str, Any]]:
    """Las filas como diccionarios, con las primeras `rellenar` columnas rellenas hacia abajo y, si
    `largo_desde`, en formato largo (columnas desde esa posición, 1-indexada → `columna`, `valor`).

    Con `filas_encabezado=2` (anexos del DANE: el año arriba, combinado, y debajo Total/Cabeceras o
    los meses), cada columna se llama «arriba abajo» y en formato largo sale `columna` (arriba) y
    `subcolumna` (abajo)."""
    nombres, partes, ultima = _encabezado(filas, fila_encabezado, filas_encabezado)
    previos: list[Any] = [None] * rellenar
    salida: list[dict[str, Any]] = []
    for fila in filas[ultima:]:
        valores = list(fila) + [None] * (len(nombres) - len(fila))
        if all(_vacio(v) for v in valores):
            continue
        for i in range(rellenar):
            if _vacio(valores[i]):
                valores[i] = previos[i]
            previos[i] = valores[i]
        registro = dict(zip(nombres, valores, strict=False))
        salida.extend([registro] if largo_desde is None else _a_largo(registro, nombres, largo_desde, partes))
    return salida


def _vacio(valor: Any) -> bool:
    return valor is None or _texto(valor) == ""


def _encabezado(filas: list[tuple[Any, ...]], fila_encabezado: int,
                filas_encabezado: int) -> tuple[list[str], dict[str, tuple[str, str]], int]:
    """(nombres, partes «arriba/abajo» de cada nombre, última fila del encabezado)."""
    if filas_encabezado not in (1, 2):
        raise ValueError("filas_encabezado es 1 o 2")
    ultima = fila_encabezado + filas_encabezado - 1
    if fila_encabezado < 1 or ultima > len(filas):
        raise ValueError(f"el encabezado queda fuera del archivo (tiene {len(filas)} filas)")
    primera, abajo = filas[fila_encabezado - 1], filas[ultima - 1]
    ancho = max(len(abajo), len(primera))
    if filas_encabezado == 1:
        return _encabezados(tuple(primera) + (None,) * (ancho - len(primera))), {}, ultima
    arriba = _arriba(primera, ancho)
    bajos = [_texto(abajo[i]) if i < len(abajo) else "" for i in range(ancho)]
    juntos = tuple(" ".join(p for p in (a, b) if p) or None for a, b in zip(arriba, bajos, strict=True))
    nombres = _encabezados(juntos)
    return nombres, dict(zip(nombres, zip(arriba, bajos, strict=True), strict=True)), ultima


def _a_largo(registro: dict[str, Any], nombres: list[str], largo_desde: int,
             partes: dict[str, tuple[str, str]]) -> list[dict[str, Any]]:
    ids = {n: registro[n] for n in nombres[: largo_desde - 1]}
    salida = []
    for n in nombres[largo_desde - 1:]:
        if _vacio(registro.get(n)):
            continue
        if n in partes:
            salida.append({**ids, "columna": partes[n][0], "subcolumna": partes[n][1], "valor": registro[n]})
        else:
            salida.append({**ids, "columna": n, "valor": registro[n]})
    return salida


def filtrar(registros: list[dict[str, Any]], filtros: dict[str, Any] | None) -> list[dict[str, Any]]:
    if not filtros:
        return registros
    columnas = set().union(*(r.keys() for r in registros)) if registros else set()
    if faltan := set(filtros) - columnas:
        raise ValueError(f"no hay columnas {sorted(faltan)}; las columnas son: {sorted(columnas)}")
    buscados = {c: _texto(v).casefold() for c, v in filtros.items()}
    return [r for r in registros if all(_texto(r.get(c)).casefold() == v for c, v in buscados.items())]
